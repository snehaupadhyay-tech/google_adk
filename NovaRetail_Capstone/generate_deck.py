from pathlib import Path

import pandas as pd
from PIL import Image

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt


# =========================================================
# 1. Project paths
# =========================================================

BASE_DIR = Path(__file__).resolve().parent
ASSETS_DIR = BASE_DIR / "assets"
ICONS_DIR = ASSETS_DIR / "Icons"
OUTPUT_DIR = BASE_DIR / "output"

DATASET_PATH = ASSETS_DIR / "novaretail_dataset.xlsx"
TEMPLATE_PATH = ASSETS_DIR / "novaretail_template.pptx"
OUTPUT_PATH = OUTPUT_DIR / "NovaRetail_QBR_Q2_2026_Final.pptx"

OUTPUT_DIR.mkdir(exist_ok=True)


# =========================================================
# 2. Load Excel data
# =========================================================

regional_df = pd.read_excel(
    DATASET_PATH,
    sheet_name="Regional_Performance",
    engine="openpyxl",
)

customer_df = pd.read_excel(
    DATASET_PATH,
    sheet_name="Customer_Experience",
    engine="openpyxl",
)

risk_df = pd.read_excel(
    DATASET_PATH,
    sheet_name="Operational_Risk",
    engine="openpyxl",
)


# =========================================================
# 3. Calculate KPIs
# =========================================================

CURRENT_QUARTER = "2026-Q2"
PREVIOUS_QUARTER = "2026-Q1"

q2_performance = regional_df[
    regional_df["Quarter"] == CURRENT_QUARTER
].copy()

q1_performance = regional_df[
    regional_df["Quarter"] == PREVIOUS_QUARTER
].copy()

q2_customer = customer_df[
    customer_df["Quarter"] == CURRENT_QUARTER
].copy()

q2_risk = risk_df[
    risk_df["Quarter"] == CURRENT_QUARTER
].copy()

total_revenue = q2_performance["Revenue_USD_M"].sum()
previous_revenue = q1_performance["Revenue_USD_M"].sum()

qoq_growth = (
    (total_revenue - previous_revenue)
    / previous_revenue
) * 100

average_nps = q2_customer["NPS"].mean()


# =========================================================
# 4. Business-content dictionary
# =========================================================

slide_content = {
    "slide1": {
        "period": "Q2 2026 | Quarterly Business Review",
        "kpi1_value": f"${total_revenue / 1000:.2f}B",
        "kpi1_label": "Q2 Revenue",
        "kpi2_value": f"+{qoq_growth:.1f}%",
        "kpi2_label": "QoQ Revenue Growth",
        "kpi3_value": f"{average_nps:.1f}",
        "kpi3_label": "Average Regional NPS",
        "summary": (
            "NovaRetail closed Q2 2026 with revenue of "
            f"${total_revenue / 1000:.2f}B, up "
            f"{qoq_growth:.1f}% from Q1. APAC delivered the "
            "strongest quarterly growth and the highest customer "
            "advocacy score. Latin America continued to contract "
            "while its customer experience and operational metrics "
            "deteriorated. APAC's capacity constraints also require "
            "attention to protect future growth."
        ),
    },

    "slide2": {
        "headline": (
            "APAC Drives Growth as Latin America Declines Further"
        ),
        "insights": [
            "APAC revenue reached $330M, increasing 8.2% QoQ.",
            "North America remained the largest market at $465M.",
            "Europe improved 4.5% QoQ but remained below its Q4 peak.",
            "Latin America declined 6.7% QoQ while returns reached 8.2%.",
        ],
    },

    "slide3": {
        "headline": (
            "APAC Sets the CX Benchmark While Latin America Deteriorates"
        ),
        "insights": [
            "APAC led customer experience with NPS 62 and CSAT of 92%.",
            "North America and Europe improved to NPS 55 and 53.",
            "Latin America NPS fell from 38 to 29 across four quarters.",
            "Delivery reached 7.5 days and support volume rose to 31K.",
        ],
    },

    "slide4": {
        "headline": (
            "Capacity Pressure and Regional Risk Demand Immediate Action"
        ),
        "insights": [
            "Latin America stockouts reached 12.5%, the highest level.",
            "Return processing in Latin America increased to 10 days.",
            "Fraud incidents rose from 18 to 30 across four quarters.",
            "APAC warehouse utilization reached 92%, constraining growth.",
        ],
    },

    "slide5": {
        "actions": [
            (
                "EXPAND APAC CAPACITY: Increase warehouse capacity "
                "and improve inventory allocation."
            ),
            (
                "STABILIZE LATIN AMERICA: Address fulfillment, "
                "returns and support workload."
            ),
            (
                "STRENGTHEN FRAUD CONTROLS: Prioritize preventive "
                "transaction monitoring."
            ),
            (
                "REPLICATE CX PRACTICES: Apply APAC delivery and "
                "support practices in weaker markets."
            ),
        ],
        "closing": (
            "Protect APAC's growth while making Latin America's "
            "recovery the immediate leadership priority."
        ),
    },
}


# =========================================================
# 5. Shape helper functions
# =========================================================

def get_shape(slide, shape_name):
    """Find a slide shape using its template name."""

    for shape in slide.shapes:
        if shape.name == shape_name:
            return shape

    available_names = [
        shape.name
        for shape in slide.shapes
    ]

    raise KeyError(
        f"Shape '{shape_name}' not found. "
        f"Available shapes: {available_names}"
    )


def remove_shape(shape):
    """Remove a PowerPoint shape from its slide."""

    shape_element = shape._element
    shape_element.getparent().remove(shape_element)


def clear_instruction_shape(slide, shape_name):
    """Remove an instruction-label shape if it exists."""

    for shape in list(slide.shapes):
        if shape.name == shape_name:
            remove_shape(shape)
            return


def set_text(
    shape,
    text,
    font_size=18,
    bold=False,
    color="1F2937",
    alignment=PP_ALIGN.LEFT,
):
    """Replace and format text in a text shape."""

    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True

    paragraph = text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = alignment

    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)


def set_bullets(shape, bullets, font_size=14):
    """Populate a shape with professionally formatted bullets."""

    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True

    text_frame.margin_left = 100000
    text_frame.margin_right = 100000
    text_frame.margin_top = 80000
    text_frame.margin_bottom = 80000

    for index, bullet_text in enumerate(bullets):

        if index == 0:
            paragraph = text_frame.paragraphs[0]
        else:
            paragraph = text_frame.add_paragraph()

        paragraph.text = f"• {bullet_text}"
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = RGBColor.from_string("374151")
        paragraph.space_after = Pt(12)
        paragraph.line_spacing = 1.1

def set_action_bullets(shape, actions):
    """Add formatted action recommendations."""

    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True

    text_frame.margin_left = 100000
    text_frame.margin_right = 100000
    text_frame.margin_top = 100000
    text_frame.margin_bottom = 100000

    for index, action in enumerate(actions):

        heading, description = action.split(":", maxsplit=1)

        if index == 0:
            paragraph = text_frame.paragraphs[0]
        else:
            paragraph = text_frame.add_paragraph()

        paragraph.level = 0
        paragraph.space_after = Pt(16)
        paragraph.line_spacing = 1.1

        bullet_run = paragraph.add_run()
        bullet_run.text = "• "
        bullet_run.font.name = "Aptos"
        bullet_run.font.size = Pt(15)
        bullet_run.font.bold = True
        bullet_run.font.color.rgb = RGBColor.from_string("0F766E")

        heading_run = paragraph.add_run()
        heading_run.text = heading.title() + ": "
        heading_run.font.name = "Aptos"
        heading_run.font.size = Pt(15)
        heading_run.font.bold = True
        heading_run.font.color.rgb = RGBColor.from_string("0F766E")

        description_run = paragraph.add_run()
        description_run.text = description.strip()
        description_run.font.name = "Aptos"
        description_run.font.size = Pt(15)
        description_run.font.color.rgb = RGBColor.from_string("374151")

# =========================================================
# 6. Icon helper
# =========================================================

def replace_with_icon(slide, placeholder_name, icon_path):
    """
    Remove the icon placeholder and insert the icon without
    stretching its original aspect ratio.
    """

    placeholder = get_shape(slide, placeholder_name)

    box_left = placeholder.left
    box_top = placeholder.top
    box_width = placeholder.width
    box_height = placeholder.height

    remove_shape(placeholder)

    with Image.open(icon_path) as image:
        image_width, image_height = image.size

    image_ratio = image_width / image_height
    box_ratio = box_width / box_height

    if image_ratio > box_ratio:
        picture_width = box_width
        picture_height = int(box_width / image_ratio)
    else:
        picture_height = box_height
        picture_width = int(box_height * image_ratio)

    picture_left = box_left + int(
        (box_width - picture_width) / 2
    )

    picture_top = box_top + int(
        (box_height - picture_height) / 2
    )

    slide.shapes.add_picture(
        str(icon_path),
        picture_left,
        picture_top,
        picture_width,
        picture_height,
    )

def apply_chart_formatting(chart, chart_type):
    """Apply consistent executive-deck formatting to charts."""

    palette = [
        "164E63",
        "0F766E",
        "14B8A6",
        "F59E0B",
    ]

    chart.chart_title.text_frame.paragraphs[0].font.name = "Aptos"
    chart.chart_title.text_frame.paragraphs[0].font.size = Pt(15)
    chart.chart_title.text_frame.paragraphs[0].font.bold = True
    chart.chart_title.text_frame.paragraphs[0].font.color.rgb = (
        RGBColor.from_string("1F2937")
    )

    chart.category_axis.tick_labels.font.name = "Aptos"
    chart.category_axis.tick_labels.font.size = Pt(10)
    chart.category_axis.tick_labels.font.color.rgb = (
        RGBColor.from_string("4B5563")
    )

    chart.value_axis.tick_labels.font.name = "Aptos"
    chart.value_axis.tick_labels.font.size = Pt(10)
    chart.value_axis.tick_labels.font.color.rgb = (
        RGBColor.from_string("4B5563")
    )

    chart.value_axis.format.line.color.rgb = (
        RGBColor.from_string("D1D5DB")
    )

    if chart_type == "column":
        for index, series in enumerate(chart.series):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = (
                RGBColor.from_string(
                    palette[index % len(palette)]
                )
            )

            series.format.line.color.rgb = (
                RGBColor.from_string(
                    palette[index % len(palette)]
                )
            )

    elif chart_type == "line":
        for index, series in enumerate(chart.series):
            series.format.line.color.rgb = (
                RGBColor.from_string(
                    palette[index % len(palette)]
                )
            )

            series.format.line.width = Pt(2.25)


# =========================================================
# 7. Chart helper functions
# =========================================================

def replace_with_performance_chart(slide):
    """Create a clustered-column revenue chart."""

    placeholder = get_shape(
        slide,
        "chart_performance",
    )

    left = placeholder.left
    top = placeholder.top
    width = placeholder.width
    height = placeholder.height

    remove_shape(placeholder)
    clear_instruction_shape(
        slide,
        "chart_performance_label",
    )

    quarters = [
        "2025-Q3",
        "2025-Q4",
        "2026-Q1",
        "2026-Q2",
    ]

    regions = [
        "North America",
        "Europe",
        "APAC",
        "Latin America",
    ]

    chart_data = CategoryChartData()
    chart_data.categories = regions

    for quarter in quarters:

        values = (
            regional_df[
                regional_df["Quarter"] == quarter
            ]
            .set_index("Region")
            .reindex(regions)["Revenue_USD_M"]
            .tolist()
        )

        chart_data.add_series(
            quarter,
            tuple(values),
        )

    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        left,
        top,
        width,
        height,
        chart_data,
    )

    chart = graphic_frame.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = (
        "Revenue by Region ($M)"
    )
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = True
    apply_chart_formatting(chart, "column")

def replace_with_customer_chart(slide):
    """Create a line chart showing NPS trends."""

    placeholder = get_shape(
        slide,
        "chart_customer",
    )

    left = placeholder.left
    top = placeholder.top
    width = placeholder.width
    height = placeholder.height

    remove_shape(placeholder)
    clear_instruction_shape(
        slide,
        "chart_customer_label",
    )

    quarters = [
        "2025-Q3",
        "2025-Q4",
        "2026-Q1",
        "2026-Q2",
    ]

    regions = [
        "North America",
        "Europe",
        "APAC",
        "Latin America",
    ]

    chart_data = CategoryChartData()
    chart_data.categories = quarters

    for region in regions:

        values = (
            customer_df[
                customer_df["Region"] == region
            ]
            .set_index("Quarter")
            .reindex(quarters)["NPS"]
            .tolist()
        )

        chart_data.add_series(
            region,
            tuple(values),
        )

    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        left,
        top,
        width,
        height,
        chart_data,
    )

    chart = graphic_frame.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = (
        "NPS Trend by Region"
    )
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False

    chart.value_axis.has_major_gridlines = True
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = 70
    apply_chart_formatting(chart, "line")


def replace_with_risk_chart(slide):
    """Create a Q2 regional stockout-rate chart."""

    placeholder = get_shape(
        slide,
        "chart_risk",
    )

    left = placeholder.left
    top = placeholder.top
    width = placeholder.width
    height = placeholder.height

    remove_shape(placeholder)
    clear_instruction_shape(
        slide,
        "chart_risk_label",
    )

    regions = [
        "North America",
        "Europe",
        "APAC",
        "Latin America",
    ]

    values = (
        q2_risk
        .set_index("Region")
        .reindex(regions)["Stockout_Rate_pct"]
        .tolist()
    )

    chart_data = CategoryChartData()
    chart_data.categories = regions

    chart_data.add_series(
        "Stockout Rate %",
        tuple(values),
    )

    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        left,
        top,
        width,
        height,
        chart_data,
    )

    chart = graphic_frame.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = (
        "Q2 2026 Stockout Rate by Region"
    )
    chart.has_legend = False

    chart.value_axis.has_major_gridlines = True
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = 15
    apply_chart_formatting(chart, "column")


# =========================================================
# 8. Populate slides
# =========================================================

def populate_slide_1(presentation):
    slide = presentation.slides[0]
    content = slide_content["slide1"]

    set_text(
        get_shape(slide, "subtitle_1"),
        content["period"],
        font_size=16,
        bold=True,
        color="4B5563",
    )

    set_text(
        get_shape(slide, "kpi1_value"),
        content["kpi1_value"],
        font_size=26,
        bold=True,
        color="0F766E",
        alignment=PP_ALIGN.CENTER,
    )

    set_text(
        get_shape(slide, "kpi1_label"),
        content["kpi1_label"],
        font_size=12,
        color="4B5563",
        alignment=PP_ALIGN.CENTER,
    )

    set_text(
        get_shape(slide, "kpi2_value"),
        content["kpi2_value"],
        font_size=26,
        bold=True,
        color="0F766E",
        alignment=PP_ALIGN.CENTER,
    )

    set_text(
        get_shape(slide, "kpi2_label"),
        content["kpi2_label"],
        font_size=12,
        color="4B5563",
        alignment=PP_ALIGN.CENTER,
    )

    set_text(
        get_shape(slide, "kpi3_value"),
        content["kpi3_value"],
        font_size=26,
        bold=True,
        color="0F766E",
        alignment=PP_ALIGN.CENTER,
    )

    set_text(
        get_shape(slide, "kpi3_label"),
        content["kpi3_label"],
        font_size=12,
        color="4B5563",
        alignment=PP_ALIGN.CENTER,
    )

    set_text(
        get_shape(slide, "summary_text"),
        content["summary"],
        font_size=15,
        color="374151",
    )

    clear_instruction_shape(slide, "icon_1_label")

    replace_with_icon(
        slide,
        "icon_1",
        ICONS_DIR / "globe_region.png",
    )


def populate_slide_2(presentation):
    slide = presentation.slides[1]
    content = slide_content["slide2"]

    set_text(
        get_shape(slide, "title_2"),
        content["headline"],
        font_size=23,
        bold=True,
        color="111827",
    )

    set_bullets(
        get_shape(slide, "insights_performance"),
        content["insights"],
    )

    replace_with_performance_chart(slide)

    clear_instruction_shape(slide, "icon_2_label")

    replace_with_icon(
        slide,
        "icon_2",
        ICONS_DIR / "growth_chart.png",
    )


def populate_slide_3(presentation):
    slide = presentation.slides[2]
    content = slide_content["slide3"]

    set_text(
        get_shape(slide, "title_3"),
        content["headline"],
        font_size=23,
        bold=True,
        color="111827",
    )

    set_bullets(
        get_shape(slide, "insights_customer"),
        content["insights"],
    )

    replace_with_customer_chart(slide)

    clear_instruction_shape(slide, "icon_3_label")

    replace_with_icon(
        slide,
        "icon_3",
        ICONS_DIR / "customer_star.png",
    )


def populate_slide_4(presentation):
    slide = presentation.slides[3]
    content = slide_content["slide4"]

    set_text(
        get_shape(slide, "title_4"),
        content["headline"],
        font_size=23,
        bold=True,
        color="111827",
    )

    set_bullets(
        get_shape(slide, "insights_risk"),
        content["insights"],
    )

    replace_with_risk_chart(slide)

    clear_instruction_shape(slide, "icon_4_label")

    replace_with_icon(
        slide,
        "icon_4",
        ICONS_DIR / "risk_alert.png",
    )


def populate_slide_5(presentation):
    slide = presentation.slides[4]
    content = slide_content["slide5"]

    set_action_bullets(
        get_shape(slide, "recommendations_text"),
        content["actions"],
    )

    set_text(
        get_shape(slide, "closing_statement"),
        content["closing"],
        font_size=17,
        bold=True,
        color="0F766E",
        alignment=PP_ALIGN.CENTER,
    )

    clear_instruction_shape(slide, "icon_5_label")

    replace_with_icon(
        slide,
        "icon_5",
        ICONS_DIR / "lightbulb_action.png",
    )


# =========================================================
# 9. Main workflow
# =========================================================

def main():

    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"Template not found: {TEMPLATE_PATH}"
        )

    if not DATASET_PATH.exists():
        raise FileNotFoundError(
            f"Dataset not found: {DATASET_PATH}"
        )

    presentation = Presentation(TEMPLATE_PATH)

    if len(presentation.slides) != 5:
        raise ValueError(
            f"Expected 5 slides, "
            f"found {len(presentation.slides)}."
        )

    populate_slide_1(presentation)
    print("Slide 1 completed")

    populate_slide_2(presentation)
    print("Slide 2 completed")

    populate_slide_3(presentation)
    print("Slide 3 completed")

    populate_slide_4(presentation)
    print("Slide 4 completed")

    populate_slide_5(presentation)
    print("Slide 5 completed")

    presentation.save(OUTPUT_PATH)

    print("\nFinal presentation created successfully:")
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()